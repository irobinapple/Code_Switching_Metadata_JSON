"""Tests for the rawmetadata transformation and exports."""

from __future__ import annotations

import io

import pandas as pd
from openpyxl import load_workbook

from src.constants import METADATA_COLUMNS, RAWMETADATA_COLUMNS
from src.exporters import (
    RAWMETADATA_SHEET_NAME,
    metadata_to_csv_bytes,
    rawmetadata_to_csv_bytes,
    rawmetadata_to_xlsx_bytes,
)
from src.transformers import (
    build_filename_stem,
    build_metadata,
    build_rawmetadata,
    conversation_duration_sec,
    segment_id,
)


class TestSegmentId:
    def test_zero_pad_width_three(self):
        assert segment_id(1) == "seg001"
        assert segment_id(42) == "seg042"

    def test_beyond_three_digits(self):
        assert segment_id(1234) == "seg1234"


class TestRawmetadataShape:
    def test_exact_column_order(self, vi_en_segments, vi_config, vi_speaker_map):
        frame = build_rawmetadata(vi_en_segments, vi_config, vi_speaker_map)
        assert list(frame.columns) == RAWMETADATA_COLUMNS
        assert len(frame.columns) == 34

    def test_one_row_per_segment(
        self, vi_en_segments, vi_config, vi_speaker_map
    ):
        frame = build_rawmetadata(vi_en_segments, vi_config, vi_speaker_map)
        assert len(frame) == len(vi_en_segments) == 4

    def test_turn_no_sequential(
        self, vi_en_segments, vi_config, vi_speaker_map
    ):
        frame = build_rawmetadata(vi_en_segments, vi_config, vi_speaker_map)
        assert frame["Turn_No"].tolist() == [1, 2, 3, 4]

    def test_segment_ids(self, vi_en_segments, vi_config, vi_speaker_map):
        frame = build_rawmetadata(vi_en_segments, vi_config, vi_speaker_map)
        assert frame["Segment_ID"].tolist() == [
            "seg001",
            "seg002",
            "seg003",
            "seg004",
        ]


class TestRawmetadataValues:
    def test_source_timestamp_text_retained(
        self, vi_en_segments, vi_config, vi_speaker_map
    ):
        frame = build_rawmetadata(vi_en_segments, vi_config, vi_speaker_map)
        assert frame["Start_Time_Sec"].iloc[0] == "00:00:04,040"
        assert frame["End_Time_Sec"].iloc[0] == "00:00:07,400"

    def test_defaults(self, vi_en_segments, vi_config, vi_speaker_map):
        frame = build_rawmetadata(vi_en_segments, vi_config, vi_speaker_map)
        assert (frame["Primary_Type"] == "Speech").all()
        assert (frame["Loudness_Level"] == "Normal").all()

    def test_language_derivations(
        self, vi_en_segments, vi_config, vi_speaker_map
    ):
        frame = build_rawmetadata(vi_en_segments, vi_config, vi_speaker_map)
        assert (frame["Segment_Primary_Language"] == "vi_VN").all()
        assert (frame["Segment_Languages"] == "vi_VN, en").all()
        assert (frame["Speaker_Languages"] == "vi_VN, en").all()
        assert (frame["LangPair"] == "vi-VN_English").all()

    def test_speaker_sources(
        self, vi_en_segments, vi_config, vi_speaker_map
    ):
        frame = build_rawmetadata(vi_en_segments, vi_config, vi_speaker_map)
        assert (frame["Speaker_Gender_Source"] == "Annotator").all()
        assert (frame["Speaker_Nativity_Source"] == "Annotator").all()
        # No No-Speaker rows here, so role source is empty everywhere.
        assert (frame["Speaker_Role_Source"] == "").all()

    def test_transliteration_blank_for_non_translit(
        self, vi_en_segments, vi_config, vi_speaker_map
    ):
        frame = build_rawmetadata(vi_en_segments, vi_config, vi_speaker_map)
        assert (frame["Transliteration_Text"] == "").all()

    def test_qc_notes_blank(self, vi_en_segments, vi_config, vi_speaker_map):
        frame = build_rawmetadata(vi_en_segments, vi_config, vi_speaker_map)
        assert (frame["QC_Notes"] == "").all()


class TestExports:
    def test_csv_utf8_bom_and_unicode(
        self, vi_en_segments, vi_config, vi_speaker_map
    ):
        frame = build_rawmetadata(vi_en_segments, vi_config, vi_speaker_map)
        data = rawmetadata_to_csv_bytes(frame)
        assert data.startswith(b"\xef\xbb\xbf")  # UTF-8 BOM
        text = data.decode("utf-8-sig")
        assert "Tôi rất bức bội." in text

    def test_xlsx_sheet_name_and_text_timestamps(
        self, vi_en_segments, vi_config, vi_speaker_map
    ):
        frame = build_rawmetadata(vi_en_segments, vi_config, vi_speaker_map)
        data = rawmetadata_to_xlsx_bytes(frame)
        workbook = load_workbook(io.BytesIO(data))
        assert workbook.sheetnames == [RAWMETADATA_SHEET_NAME]
        worksheet = workbook[RAWMETADATA_SHEET_NAME]

        header = [c.value for c in worksheet[1]]
        start_col = header.index("Start_Time_Sec") + 1
        # Data cell stores the source string and is Text-formatted.
        cell = worksheet.cell(row=2, column=start_col)
        assert cell.value == "00:00:04,040"
        assert cell.number_format == "@"
        assert isinstance(cell.value, str)

    def test_xlsx_unicode_content(
        self, vi_en_segments, vi_config, vi_speaker_map
    ):
        frame = build_rawmetadata(vi_en_segments, vi_config, vi_speaker_map)
        data = rawmetadata_to_xlsx_bytes(frame)
        workbook = load_workbook(io.BytesIO(data))
        worksheet = workbook[RAWMETADATA_SHEET_NAME]
        header = [c.value for c in worksheet[1]]
        content_col = header.index("Content_Text") + 1
        values = [
            worksheet.cell(row=r, column=content_col).value
            for r in range(2, worksheet.max_row + 1)
        ]
        assert any("Tôi rất bức bội." in (v or "") for v in values)


class TestMetadata:
    def test_exact_column_order(
        self, vi_en_segments, vi_config, vi_speaker_map
    ):
        frame = build_metadata(vi_en_segments, vi_config, vi_speaker_map)
        assert list(frame.columns) == METADATA_COLUMNS
        assert len(frame.columns) == 21

    def test_one_row_per_unique_speaker(
        self, vi_en_segments, vi_config, vi_speaker_map
    ):
        frame = build_metadata(vi_en_segments, vi_config, vi_speaker_map)
        assert len(frame) == 2
        assert frame["Speaker_ID"].tolist() == ["S1", "S2"]
        assert frame["SNo"].tolist() == [1, 2]

    def test_per_speaker_turn_count_not_conversation_total(
        self, vi_en_segments, vi_config, vi_speaker_map
    ):
        # SPK001 speaks turns 1 & 3, SPK002 turns 2 & 4 -> 2 each, not 4.
        frame = build_metadata(vi_en_segments, vi_config, vi_speaker_map)
        counts = dict(
            zip(frame["Speaker_ID"], frame["Number_of_Turns"])
        )
        assert counts == {"S1": 2, "S2": 2}

    def test_duration_is_conversation_level(
        self, vi_en_segments, vi_config, vi_speaker_map
    ):
        frame = build_metadata(vi_en_segments, vi_config, vi_speaker_map)
        # max end (24.8) - min start (4.04) = 20.76, same on every row.
        assert frame["Duration_Sec"].tolist() == [20.76, 20.76]

    def test_duration_numeric(
        self, vi_en_segments, vi_config, vi_speaker_map
    ):
        frame = build_metadata(vi_en_segments, vi_config, vi_speaker_map)
        assert pd.api.types.is_numeric_dtype(frame["Duration_Sec"])

    def test_filename_underscore_joined(
        self, vi_en_segments, vi_config, vi_speaker_map
    ):
        frame = build_metadata(vi_en_segments, vi_config, vi_speaker_map)
        assert (
            frame["FileName"] == "vi-VN_English_AIR_48kHz_Conv0347"
        ).all()

    def test_no_qc_notes_column(
        self, vi_en_segments, vi_config, vi_speaker_map
    ):
        frame = build_metadata(vi_en_segments, vi_config, vi_speaker_map)
        assert "QC_Notes" not in frame.columns

    def test_metadata_csv_bom_and_unicode(
        self, vi_en_segments, vi_config, vi_speaker_map
    ):
        frame = build_metadata(vi_en_segments, vi_config, vi_speaker_map)
        data = metadata_to_csv_bytes(frame)
        assert data.startswith(b"\xef\xbb\xbf")


class TestDecimalCsRatios:
    """CS ratios may be fractional, e.g. a 60.8 / 39.2 split."""

    def _decimal_config(self, vi_config):
        import dataclasses

        return dataclasses.replace(
            vi_config, cs_ratio_primary=60.8, cs_ratio_secondary=39.2
        )

    def test_rawmetadata_keeps_decimals(
        self, vi_en_segments, vi_config, vi_speaker_map
    ):
        cfg = self._decimal_config(vi_config)
        frame = build_rawmetadata(vi_en_segments, cfg, vi_speaker_map)
        assert (frame["CS_Ratio_Primary"] == 60.8).all()
        assert (frame["CS_Ratio_Secondary"] == 39.2).all()

    def test_metadata_keeps_decimals(
        self, vi_en_segments, vi_config, vi_speaker_map
    ):
        cfg = self._decimal_config(vi_config)
        frame = build_metadata(vi_en_segments, cfg, vi_speaker_map)
        assert (frame["CS_Ratio_Primary"] == 60.8).all()
        assert (frame["CS_Ratio_Secondary"] == 39.2).all()

    def test_decimal_split_totalling_100_does_not_warn(
        self, vi_en_segments, vi_config, vi_speaker_map
    ):
        from src.validators import validate_rawmetadata

        cfg = self._decimal_config(vi_config)
        frame = build_rawmetadata(vi_en_segments, cfg, vi_speaker_map)
        result = validate_rawmetadata(
            frame, cfg, vi_speaker_map, vi_en_segments
        )
        # 60.8 + 39.2 must not trip the "ratios do not total 100" warning
        # despite binary floating-point representation.
        assert not any("total" in w for w in result.warnings), result.warnings


class TestHelpers:
    def test_filename_stem(self, vi_config):
        assert build_filename_stem(vi_config) == (
            "vi-VN_English_AIR_48kHz_Conv0347"
        )

    def test_duration_helper(self, vi_en_segments):
        assert conversation_duration_sec(vi_en_segments) == 20.76

    def test_duration_empty(self):
        assert conversation_duration_sec([]) == 0.0
